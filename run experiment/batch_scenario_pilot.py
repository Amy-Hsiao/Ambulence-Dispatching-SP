"""batch_scenario_pilot.py — 情境數 (|S|) 掃描 pilot：決定正式 ablation 要用哪個 S。

為什麼要做這個
--------------
2026-09-05 的 DRO_full_matrix 用 S=5 跑完 3×3×8=72 格，結果有三個問題：

  1. Extensive 在 S=5 之下太好解（large 只有 363,248 個變數），gap 只有 2.7~2.8%，
     反而贏過階梯的第 2~7 段 —— 老師要的「B&BC 有效」看不出來。
  2. CVaR 退化：alpha=0.9 的尾端質量 = 0.1，但 S=5 時每個情境機率 0.2，
     尾端連一個情境都不到 → CVaR 直接等於最差情境，第 4 章的 CVaR 線性化形同虛設。
  3. 三個模糊集分不出來：box 的實質強度是 eps_bar × S，S=5 且 eps=0.01 時
     只有 ±5%（S=30 時是 ±30%）。三組 Extensive 目標值差距 < 0.09%。

另一邊，2026-08-17 的 CCP_count_ablation 用 S=50，Extensive 在 12 格裡
全部 `nodes=1`、gap 44~84%，但 B&BC 的 root seeding 在 large/J>=30 吃掉
96~97% 的時限，LB=-inf、nodes=0，整個垮掉。

所以 S 太小不行、太大也不行。這支程式就是去找中間那個值。

做什麼
------
固定 |J|=20（SCALE_PROFILES 已是 20），固定規模（預設只跑 large，最難的那個），
掃 S ∈ {10, 20, 30}，每個 S 只跑兩個 config：

    Extensive    整體式基準線，一條 VI 都不加、零加速     ← 要它「爛」
    FullStack    BBC+WS+RS+UC+Pareto+LBF+VI            ← 要它 <= 5%

挑選標準（Verdict 分頁會自動算）
    Extensive gap  >  EXT_MIN_GAP_PCT (預設 15%)   整體式明顯吃力
    FullStack gap  <= FS_MAX_GAP_PCT  (預設 5%)    全堆疊仍在可接受範圍
    兩者都成立的「最大 S」就是正式 ablation 該用的值。

eps_bar 怎麼跟著 S 調（--eps-mode）
----------------------------------
box 模糊集 {p : |p_s - p0_s| <= eps_bar} 的實質強度是相對擾動 eps_bar / p0_s
= eps_bar × S。若固定 eps_bar=0.01，S 從 30 換到 5 會讓強度從 ±30% 掉到 ±5%。

    scaled (預設)  固定 eps_bar × S = EPS_BOX_PRODUCT (= 0.01 × 30 = 0.30)
                   → S=10 時 eps=0.030、S=20 時 0.015、S=30 時 0.010
                   與 Thesis_Draft_(2026.08.10) 的校準情境一致，跨 S 可比。
    fixed          一律用 config.DRO_EPSILON_BOX，不隨 S 調（僅供對照）

註：risk_core 要求 box 的 eps_bar <= min_s p0_s = 1/S。
    scaled 模式恆滿足（0.30/S <= 1/S ⟺ 0.30 <= 1）。程式仍會實際驗一次。
    --model mcvar 時沒有模糊集，此選項不生效。

執行
----
    python "run experiment/batch_scenario_pilot.py"
    python "run experiment/batch_scenario_pilot.py" --scenarios 10,20,30 --scales large
    python "run experiment/batch_scenario_pilot.py" --model dro_box
    python "run experiment/batch_scenario_pilot.py" --time-limit 600 --scenarios 10
    python "run experiment/batch_scenario_pilot.py" --dry-run     # 不需 Gurobi，驗流程用

輸出（experiment result/）
    scenario_pilot_<timestamp>.xlsx
        Verdict   ★ 先看這張：每個 S 的兩個 gap、是否通過、最後的建議
        Summary   UB / LB / Time / Gap(%) / Nodes（每個 S × 每個 config 一列）
        Detail    切割數、root seeding 診斷、CVaR 尾端、模型規模等完整欄位
        Config    本次每一項設定，供重現
    scenario_pilot_raw_<timestamp>.csv
        每跑完一格就立刻重寫，中途掛掉不會白跑。

續跑
    把上面那個 raw csv 的檔名丟給 --resume-from，已完成且狀態正常的格子會跳過。
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
import time
import traceback
from datetime import datetime, timedelta
from pathlib import Path

for _st in (sys.stdout, sys.stderr):
    try:
        _st.reconfigure(encoding="utf-8", errors="replace")
    except Exception:                                          # noqa: BLE001
        pass

ROOT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT_DIR / "model core"))
sys.path.insert(0, str(ROOT_DIR / "model portal"))
RESULT_DIR = ROOT_DIR / "experiment result"

import gurobipy as gp                                          # noqa: E402

import config                                                  # noqa: E402
import extensive_form_core as model_core                       # noqa: E402
import lshaped_core                                            # noqa: E402
import risk_core                                               # noqa: E402
# Extensive 的風險目標式直接重用 portal 既有入口，確保與 batch_dro_full_matrix
# 逐字相同 —— 自己重寫一份會有對不上的風險。
import extensive_dro                                           # noqa: E402

RESULT_PREFIX = "scenario_pilot"

# ── 掃描與判定的預設值 ───────────────────────────────────────────────── #
DEFAULT_SCENARIOS = (10, 20, 30)
DEFAULT_SCALES = ("large",)
DEFAULT_TIME_LIMIT = 7200.0
DEFAULT_MIP_GAP = 0.01

EXT_MIN_GAP_PCT = 15.0      # Extensive gap 要 > 這個值，才算「整體式明顯吃力」
FS_MAX_GAP_PCT = 5.0        # 全堆疊 gap 要 <= 這個值

# eps_bar × S 的不變量：Thesis_Draft_(2026.08.10) 的校準情境是 S=30、eps=0.01
EPS_BOX_REF_S = 30
EPS_BOX_REF_VALUE = 0.01
EPS_BOX_PRODUCT = EPS_BOX_REF_VALUE * EPS_BOX_REF_S            # = 0.30

_ALL_VI_OFF = {"all": False}
_ALL_VI_ON = {"all": True}

# 這兩段就是 batch_dro_full_matrix.py LADDER 的第 1 段與第 8 段，逐項對齊。
CONFIGS = (
    ("Extensive", dict(kind="ext")),
    ("FullStack", dict(kind="bbc", ws=True, rs=None, uc=True,
                       pareto=True, lbf=True, vi=True)),
)

SUCCESS_STATUS = {"OPTIMAL", "TIME_LIMIT", "SUBOPTIMAL", "INTERRUPTED",
                  "NODE_LIMIT", "SOLUTION_LIMIT", "ITERATION_LIMIT"}
STATUS_NAME = {1: "LOADED", 2: "OPTIMAL", 3: "INFEASIBLE", 4: "INF_OR_UNBD",
               5: "UNBOUNDED", 6: "CUTOFF", 7: "ITERATION_LIMIT", 8: "NODE_LIMIT",
               9: "TIME_LIMIT", 10: "SOLUTION_LIMIT", 11: "INTERRUPTED",
               12: "NUMERIC", 13: "SUBOPTIMAL", 14: "INPROGRESS",
               15: "USER_OBJ_LIMIT", 16: "WORK_LIMIT", 17: "MEM_LIMIT"}

CSV_FIELDS = [
    "scale", "n_scen", "config", "model", "n_I", "n_J", "n_H", "n_T", "n_S",
    "risk_alpha", "risk_lambda", "eps_mode", "ambiguity_scope",
    "cvar_tail_scenarios", "dro_relative_strength",
    "UB", "LB", "Gap_pct", "Time", "Nodes", "status",
    "solver_secs", "build_secs", "wall_secs",
    "model_vars", "model_constrs", "model_qconstrs",
    "iterations", "cuts_added", "seed_cuts", "user_cuts", "lazy_cuts",
    "root_seed_iters_done", "root_seed_lb", "root_seed_stop",
    "root_seed_time", "root_seed_time_budget", "root_seed_pct_of_limit",
    "root_cut_rounds_done", "oracle_solves", "cache_hits", "callback_time",
    "callback_pct_of_wall", "vi_flags", "ws_effective",
    "sum_X", "sum_V", "sum_U", "sum_Y",
    "started_at", "finished_at", "sanity", "error",
]


def sname(st) -> str:
    try:
        return STATUS_NAME.get(int(st), f"STATUS_{st}")
    except (TypeError, ValueError):
        return str(st)


# ===================================================================== #
# eps_bar 隨 S 縮放                                                      #
# ===================================================================== #
def resolve_epsilon_box(n_scen: int, eps_mode: str) -> float:
    """回傳這個 S 該用的 box eps_bar。scaled 模式固定 eps_bar × S。"""
    if eps_mode == "fixed":
        return float(getattr(config, "DRO_EPSILON_BOX", 0.01))
    if eps_mode != "scaled":
        raise ValueError(f"未知的 eps_mode: {eps_mode!r}")
    if n_scen <= 0:
        raise ValueError("n_scen 必須 >= 1")
    return EPS_BOX_PRODUCT / float(n_scen)


def make_risk_cfg_for(model_type: str, n_scen: int, eps_mode: str):
    """mcvar 直接用 config 預設；dro_box 依 eps_mode 決定 eps_bar。"""
    if model_type == "mcvar":
        return risk_core.make_risk_cfg("mcvar")
    if model_type == "dro_box":
        return risk_core.make_risk_cfg(
            "dro_box", epsilon_box=resolve_epsilon_box(n_scen, eps_mode))
    raise ValueError(f"未支援的 model: {model_type!r}（只支援 mcvar / dro_box）")


# ===================================================================== #
# instance 快取：一次只留一個 (scale, S)，避免大 instance 疊在記憶體裡      #
# ===================================================================== #
_INSTANCE_CACHE: dict[tuple[str, int], dict] = {}


def get_instance(scale: str, n_scen: int) -> dict:
    """取得指定規模與情境數的 instance。

    關鍵：config.generate_data() 在函式內部讀模組層級的 SCENARIOS，
    所以必須「先改 config.SCENARIOS 再呼叫 generate_data」才會真的
    產生 n_scen 個情境。既有的 S_all[:n] 只能往下砍、不能往上加。
    情境是依 index 決定亂數種子，故 S=5 ⊂ S=10 ⊂ S=30（巢狀、跨 S 可比）。
    """
    key = (scale, n_scen)
    if key not in _INSTANCE_CACHE:
        _INSTANCE_CACHE.clear()                                # 只留最新的一個
        prev = getattr(config, "SCENARIOS", None)
        t0 = time.time()
        try:
            config.SCENARIOS = int(n_scen)
            inst = config.generate_data(scale=scale)
        finally:
            if prev is not None:
                config.SCENARIOS = prev
        got = len(inst["sets"]["S"])
        if got != n_scen:
            raise RuntimeError(
                f"要求 {n_scen} 個情境，generate_data 只給了 {got} 個 —— "
                "請檢查 config.generate_data 是否仍在函式內讀 SCENARIOS。")
        _INSTANCE_CACHE[key] = inst
        print(f"   （生成 {scale} / S={n_scen} instance：{time.time() - t0:.1f}s）",
              flush=True)
    return _INSTANCE_CACHE[key]


def norm_probs_of(instance: dict, S_sel: list[str]) -> dict[str, float]:
    sd = instance["scenario_data"]
    raw = {s: sd["probability"][s] for s in S_sel}
    tot = sum(raw.values())
    return {s: p / tot for s, p in raw.items()}


# ===================================================================== #
# 兩條求解路徑                                                           #
# ===================================================================== #
def solve_extensive(instance, S_sel, risk_cfg, time_limit, mip_gap) -> dict:
    """整體式：建完整模型 + 風險目標式，直接丟給 Gurobi。零加速、零 VI。"""
    sets = instance["sets"]
    p0 = norm_probs_of(instance, S_sel)
    sub_sd = {k: {s: instance["scenario_data"][k][s] for s in S_sel} for k in
              ("demand", "road_availability_ij", "road_availability_jh",
               "hospital_receiving_capacity")}
    m = None
    t_build = time.time()
    try:
        m, v = model_core.build_gurobi_model(
            sets["I"], sets["J"], sets["H"], sets["L"], sets["L_transfer"],
            sets["T"], S_sel, instance["deterministic_parameters"], sub_sd, p0,
            instance["road_capacity"]["cap_ij"], instance["road_capacity"]["cap_jh"],
            instance["transport_cost"]["cost_ij"], instance["transport_cost"]["cost_jh"],
            model_name="Extensive_Pilot", time_limit=time_limit, mip_gap=mip_gap,
            vi_cfg=_ALL_VI_OFF,          # Extensive 是基準線，一條 VI 都不加
        )
        extensive_dro._apply_risk_objective(m, v, S_sel, p0, risk_cfg)
        build_secs = time.time() - t_build
        m.setParam("OutputFlag", 1)
        m.setParam("DisplayInterval",
                   int(getattr(config, "BENDERS_DISPLAY_INTERVAL", 30)))
        m.optimize()
        has = m.SolCount > 0
        out = dict(
            UB=float(m.ObjVal) if has else None,
            Gap_pct=float(m.MIPGap) * 100.0 if has else None,
            solver_secs=float(m.Runtime),
            build_secs=build_secs,
            status=sname(m.status),
            model_vars=int(m.NumVars), model_constrs=int(m.NumConstrs),
            model_qconstrs=int(m.NumQConstrs),
        )
        for key, attr in (("LB", "ObjBound"), ("Nodes", "NodeCount")):
            try:
                out[key] = float(getattr(m, attr))
            except (gp.GurobiError, AttributeError):
                out[key] = None
        return out
    finally:
        if m is not None:
            m.dispose()


def solve_full_stack(instance, S_sel, risk_cfg, cfg, time_limit, mip_gap) -> dict:
    """全堆疊：BBC + WS + RS + UC + Pareto + LBF + VI。"""
    ws = cfg["ws"] and ("deterministic_data" in instance)
    rs = cfg["rs"]
    if rs is None:
        rs = int(getattr(config, "BENDERS_ROOT_SEED_ITERS", 300))
    uc_rounds = int(getattr(config, "BENDERS_ROOT_CUT_ROUNDS", 15))
    res = master = None
    try:
        res = lshaped_core.solve_bbc(
            instance, S_sel,
            time_limit=time_limit, mip_gap=mip_gap, risk_cfg=risk_cfg,
            multi_cut=True,                                    # 風險模型必須 multi-cut
            ev_warm_start=ws,                                  # WS
            root_seed_iters=rs,                                # RS
            use_user_cuts=cfg["uc"],                           # UC
            root_cut_rounds=(uc_rounds if cfg["uc"] else 0),
            pareto_enabled=cfg["pareto"],                      # Pareto
            lbf_enabled=cfg["lbf"],                            # LBF
            # 明確傳 {"all": True}，不靠 config.VI_ENABLED —— 否則有人改了
            # config，這一段就會在標籤寫著 VI 的情況下沒有 VI。
            vi_cfg=(_ALL_VI_ON if cfg["vi"] else _ALL_VI_OFF),
            verbose=True,
        )
        fs = res.get("first_stage") or {}
        master = res.pop("master", None)
        res.pop("vars", None)
        # solve_bbc 在沒有 incumbent 時把 nodes 記成 0，會讓最難的格子看起來
        # 像「根節點就解完」。直接讀 master 的真實節點數。
        nodes = res.get("nodes")
        nq = None
        if master is not None:
            try:
                nodes = float(master.NodeCount)
                nq = int(master.NumQConstrs)
            except (gp.GurobiError, AttributeError):
                pass
        return dict(
            UB=res.get("best_ub"), LB=res.get("best_lb"),
            solver_secs=res.get("runtime"),
            Gap_pct=res.get("gap_pct"), Nodes=nodes,
            status=res.get("status"),
            ws_effective=ws, model_qconstrs=nq,
            iterations=res.get("iterations"), cuts_added=res.get("cuts_added"),
            seed_cuts=res.get("seed_cuts_added"), user_cuts=res.get("user_cuts_added"),
            lazy_cuts=res.get("lazy_cuts_added"),
            root_seed_iters_done=res.get("root_seed_iters_done"),
            root_seed_lb=res.get("root_seed_lb"),
            root_seed_stop=res.get("root_seed_stop_reason"),
            root_seed_time=res.get("root_seed_time"),
            root_seed_time_budget=res.get("root_seed_time_budget"),
            root_cut_rounds_done=res.get("root_cut_rounds_done"),
            oracle_solves=res.get("oracle_solves"), cache_hits=res.get("cache_hits"),
            callback_time=res.get("callback_time"),
            vi_flags=json.dumps(res.get("vi_flags") or {}, ensure_ascii=False),
            sum_X=sum(fs.get("X", {}).values()) if fs else None,
            sum_V=sum(fs.get("V", {}).values()) if fs else None,
            sum_U=sum(fs.get("U", {}).values()) if fs else None,
            sum_Y=sum(fs.get("Y", {}).values()) if fs else None,
        )
    finally:
        if master is not None:
            try:
                master.dispose()
            except Exception:                                  # noqa: BLE001
                pass


# ===================================================================== #
# 假求解器（--dry-run）：不碰 Gurobi，只驗流程與輸出                        #
# ===================================================================== #
def fake_solve(config_name: str, n_scen: int, time_limit: float) -> dict:
    """數字是編的，只用來確認欄位、CSV、xlsx、Verdict 判定都跑得通。"""
    time.sleep(0.01)
    if config_name == "Extensive":
        gap = 2.8 + 0.9 * n_scen                               # S 越大整體式越爛
        return dict(UB=2.6e7, LB=2.6e7 * (1 - gap / 100), Gap_pct=gap,
                    Nodes=max(1.0, 40.0 - n_scen), status="TIME_LIMIT",
                    solver_secs=time_limit, build_secs=3.0,
                    model_vars=72644 * n_scen // 1, model_constrs=29200 * n_scen,
                    model_qconstrs=0)
    gap = 2.0 + 0.10 * n_scen                                  # 全堆疊緩慢變差
    return dict(UB=2.63e7, LB=2.63e7 * (1 - gap / 100), Gap_pct=gap,
                Nodes=4800.0, status="TIME_LIMIT", solver_secs=time_limit,
                build_secs=0.0, ws_effective=True, model_qconstrs=0,
                iterations=1.6e7, cuts_added=999, seed_cuts=130, user_cuts=150,
                lazy_cuts=719, root_seed_iters_done=14, root_seed_lb=2.49e7,
                root_seed_stop="lb_rel_improve_below_0.000500_for_10_rounds",
                root_seed_time=239.0, root_seed_time_budget=0.15 * time_limit,
                root_cut_rounds_done=15, oracle_solves=1075, cache_hits=3,
                callback_time=363.0, vi_flags=json.dumps({f"VI-{i}": True
                                                          for i in range(1, 9)}),
                sum_X=7.0, sum_V=279.0, sum_U=97.0, sum_Y=5584.0)


# ===================================================================== #
# 自動健檢                                                               #
# ===================================================================== #
def sanity_flags(row: dict, time_limit: float) -> str:
    f = []
    if row.get("status") in SUCCESS_STATUS:
        if row.get("LB") is None:
            f.append("無下界")
        if row.get("UB") is None:
            f.append("無可行解")
        nd = row.get("Nodes")
        if isinstance(nd, (int, float)) and nd == 0 and row.get("config") != "Extensive":
            f.append("節點數為0（B&C沒開始）")
        rst = row.get("root_seed_time")
        if isinstance(rst, (int, float)) and rst > 0.25 * time_limit:
            f.append(f"root_seeding佔{rst / time_limit * 100:.0f}%時限")
        if row.get("config") == "FullStack":
            if row.get("ws_effective") is False:
                f.append("暖啟動未生效")
            try:
                d = json.loads(row.get("vi_flags") or "{}")
                if d and not all(d.values()):
                    f.append("VI未全開")
            except (ValueError, TypeError):
                pass
    return "; ".join(f)


# ===================================================================== #
# 單一格                                                                 #
# ===================================================================== #
def run_cell(scale: str, n_scen: int, config_name: str, cfg: dict,
             model_type: str, eps_mode: str,
             time_limit: float, mip_gap: float, dry_run: bool) -> dict:
    t0 = time.time()
    row = dict.fromkeys(CSV_FIELDS)
    row.update(scale=scale, n_scen=n_scen, config=config_name, model=model_type,
               eps_mode=eps_mode, started_at=datetime.now().isoformat(timespec="seconds"),
               status="NOT_RUN")
    print(f"\n{'=' * 72}\n  {scale} | S={n_scen} | {config_name} | {model_type} | "
          f"時限 {time_limit / 3600:.2f} 小時\n{'=' * 72}", flush=True)
    try:
        risk_cfg = make_risk_cfg_for(model_type, n_scen, eps_mode)
        row.update(risk_alpha=risk_cfg["alpha"], risk_lambda=risk_cfg["lambda"],
                   ambiguity_scope=risk_core.risk_scope(risk_cfg))
        # CVaR 尾端有幾個情境：(1-alpha)/p0_s = (1-alpha)*S。< 1 代表 CVaR 退化成 max。
        row["cvar_tail_scenarios"] = round((1.0 - risk_cfg["alpha"]) * n_scen, 3)
        scope = row["ambiguity_scope"]
        row["dro_relative_strength"] = (round(scope * n_scen, 4)
                                        if scope is not None else None)

        if dry_run:
            row.update(n_I=0, n_J=0, n_H=0, n_T=0, n_S=n_scen)
            row.update(fake_solve(config_name, n_scen, time_limit))
        else:
            instance = get_instance(scale, n_scen)
            sets = instance["sets"]
            S_sel = list(sets["S"])
            row.update(n_I=len(sets["I"]), n_J=len(sets["J"]), n_H=len(sets["H"]),
                       n_T=len(sets["T"]), n_S=len(S_sel))
            # box 的 eps_bar <= min_s p0_s，不合格會在這裡就擋下來
            risk_core.validate_risk_cfg_for_probs(
                risk_cfg, norm_probs_of(instance, S_sel))
            if cfg["kind"] == "ext":
                row.update(solve_extensive(instance, S_sel, risk_cfg,
                                           time_limit, mip_gap))
            else:
                row.update(solve_full_stack(instance, S_sel, risk_cfg, cfg,
                                            time_limit, mip_gap))
    except gp.GurobiError as exc:
        row.update(status="GurobiError", error=str(exc))
        print(f"\n   [!] Gurobi 錯誤：{exc}", flush=True)
    except MemoryError:
        row.update(status="MemoryError", error="記憶體不足")
        print("\n   [!] 記憶體不足", flush=True)
    except Exception as exc:                                   # noqa: BLE001
        row.update(status="Exception", error=f"{type(exc).__name__}: {exc}")
        traceback.print_exc()

    # Time 一律用整格牆鐘時間（含建模），Extensive 與 BBC 才可比。
    row["Time"] = time.time() - t0
    row["wall_secs"] = row["Time"]
    row["finished_at"] = datetime.now().isoformat(timespec="seconds")
    rst, cbt = row.get("root_seed_time"), row.get("callback_time")
    if isinstance(rst, (int, float)) and time_limit > 0:
        row["root_seed_pct_of_limit"] = round(rst / time_limit * 100, 2)
    if isinstance(cbt, (int, float)) and row["wall_secs"] > 0:
        row["callback_pct_of_wall"] = round(cbt / row["wall_secs"] * 100, 2)
    row["sanity"] = sanity_flags(row, time_limit)

    g = row.get("Gap_pct")
    print(f"\n   → status={row['status']}  gap="
          f"{f'{g:.3f}%' if isinstance(g, (int, float)) else 'NA'}  "
          f"用時 {row['Time'] / 60:.1f} 分"
          + (f"  [健檢] {row['sanity']}" if row["sanity"] else ""), flush=True)
    return row


# ===================================================================== #
# Verdict：把挑選標準算成一張表                                            #
# ===================================================================== #
def build_verdict(rows: list[dict], ext_min_gap: float, fs_max_gap: float) -> list[dict]:
    by: dict[tuple[str, int], dict[str, dict]] = {}
    for r in rows:
        by.setdefault((r["scale"], int(r["n_scen"])), {})[r["config"]] = r
    out = []
    for (scale, n_scen) in sorted(by, key=lambda k: (k[0], k[1])):
        ext = by[(scale, n_scen)].get("Extensive") or {}
        fst = by[(scale, n_scen)].get("FullStack") or {}
        eg, fg = ext.get("Gap_pct"), fst.get("Gap_pct")
        ext_ok = isinstance(eg, (int, float)) and eg > ext_min_gap
        fs_ok = isinstance(fg, (int, float)) and fg <= fs_max_gap
        tail = fst.get("cvar_tail_scenarios") or ext.get("cvar_tail_scenarios")
        note = []
        if isinstance(tail, (int, float)) and tail < 1:
            note.append(f"CVaR尾端只有{tail}個情境（退化成最差情境）")
        if not isinstance(eg, (int, float)):
            note.append("Extensive無gap")
        if not isinstance(fg, (int, float)):
            note.append("FullStack無gap")
        for r, lbl in ((ext, "Ext"), (fst, "FS")):
            if r.get("sanity"):
                note.append(f"{lbl}健檢:{r['sanity']}")
        out.append({
            "scale": scale, "n_scen": n_scen,
            "Extensive_gap_pct": eg, "FullStack_gap_pct": fg,
            "Extensive_nodes": ext.get("Nodes"), "FullStack_nodes": fst.get("Nodes"),
            "cvar_tail_scenarios": tail,
            "dro_relative_strength": fst.get("dro_relative_strength"),
            "root_seed_pct_of_limit": fst.get("root_seed_pct_of_limit"),
            f"Extensive>{ext_min_gap:g}%": "PASS" if ext_ok else "FAIL",
            f"FullStack<={fs_max_gap:g}%": "PASS" if fs_ok else "FAIL",
            "verdict": "OK" if (ext_ok and fs_ok) else "NO",
            "note": "; ".join(note),
        })
    return out


def recommend(verdict_rows: list[dict]) -> str:
    ok = [v for v in verdict_rows if v["verdict"] == "OK"]
    if not ok:
        return ("沒有任何 S 同時滿足兩個條件。若 Extensive 都太好解 → 再往上加 S；"
                "若 FullStack 都超過門檻 → 先降 S，或檢查 root seeding 是否又吃光時間。")
    best = max(ok, key=lambda v: v["n_scen"])
    return (f"建議正式 ablation 使用 S = {best['n_scen']}"
            f"（{best['scale']}：Extensive gap {best['Extensive_gap_pct']:.2f}%、"
            f"FullStack gap {best['FullStack_gap_pct']:.2f}%）")


# ===================================================================== #
# 輸出                                                                   #
# ===================================================================== #
def write_csv(path: Path, rows: list[dict]) -> None:
    tmp = path.with_suffix(".csv.tmp")
    with tmp.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=CSV_FIELDS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    tmp.replace(path)                                          # 原子性取代


def write_xlsx(path: Path, rows: list[dict], verdict_rows: list[dict],
               settings: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Verdict"
    if verdict_rows:
        heads = list(verdict_rows[0])
        ws.append(heads)
        for c in ws[1]:
            c.font = bold
        for v in verdict_rows:
            ws.append([v.get(h) for h in heads])
        ws.append([])
        ws.append(["建議", recommend(verdict_rows)])
        ws.cell(row=ws.max_row, column=1).font = bold
    else:
        ws.append(["（沒有任何結果）"])

    sm = wb.create_sheet("Summary")
    sm_cols = ["scale", "n_scen", "config", "model", "n_I", "n_J", "n_H", "n_S",
               "UB", "LB", "Time", "Gap_pct", "Nodes", "status", "sanity"]
    sm.append(sm_cols)
    for c in sm[1]:
        c.font = bold
    for r in rows:
        sm.append([r.get(k) for k in sm_cols])

    dt = wb.create_sheet("Detail")
    dt.append(CSV_FIELDS)
    for c in dt[1]:
        c.font = bold
    for r in rows:
        dt.append([r.get(k) for k in CSV_FIELDS])

    cf = wb.create_sheet("Config")
    cf.append(["設定", "值"])
    for c in cf[1]:
        c.font = bold
    for k, v in settings.items():
        cf.append([k, v])

    for sheet in (ws, sm, dt, cf):
        for col in sheet.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None),
                        default=8)
            sheet.column_dimensions[col[0].column_letter].width = min(42, max(10, width + 2))

    tmp = path.with_name(path.stem + ".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)


# ===================================================================== #
# CLI                                                                    #
# ===================================================================== #
def parse_args(argv=None):
    ap = argparse.ArgumentParser(
        description="情境數 (|S|) 掃描 pilot：決定正式 ablation 要用哪個 S。",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    ap.add_argument("--scenarios", default=",".join(str(s) for s in DEFAULT_SCENARIOS),
                    help="要掃的情境數，逗號分隔")
    ap.add_argument("--scales", default=",".join(DEFAULT_SCALES),
                    help="規模，逗號分隔（small/medium/large）")
    ap.add_argument("--model", default="mcvar", choices=("mcvar", "dro_box"),
                    help="mcvar = SP+MCVaR；dro_box = MCVaR + box DRO")
    ap.add_argument("--eps-mode", default="scaled", choices=("scaled", "fixed"),
                    help="dro_box 的 eps_bar：scaled 固定 eps×S，fixed 用 config 值")
    ap.add_argument("--time-limit", type=float, default=DEFAULT_TIME_LIMIT,
                    help="每格時限（秒）")
    ap.add_argument("--mip-gap", type=float, default=DEFAULT_MIP_GAP,
                    help="相對 MIP gap 提早結束門檻")
    ap.add_argument("--ext-min-gap", type=float, default=EXT_MIN_GAP_PCT,
                    help="判定用：Extensive gap 要大於此值（%%）")
    ap.add_argument("--fs-max-gap", type=float, default=FS_MAX_GAP_PCT,
                    help="判定用：FullStack gap 要小於等於此值（%%）")
    ap.add_argument("--configs", default="Extensive,FullStack",
                    help="要跑哪些 config，逗號分隔")
    ap.add_argument("--resume-from", default=None,
                    help="先前的 raw csv 路徑；已完成且狀態正常的格子會跳過")
    ap.add_argument("--dry-run", action="store_true",
                    help="不呼叫 Gurobi，用假結果驗整條流程與輸出")
    a = ap.parse_args(argv)

    try:
        a.scenario_list = [int(x) for x in a.scenarios.split(",") if x.strip()]
    except ValueError:
        ap.error(f"--scenarios 只能是整數：{a.scenarios!r}")
    if not a.scenario_list:
        ap.error("--scenarios 不能是空的")
    if any(s < 1 for s in a.scenario_list):
        ap.error("--scenarios 每個值至少要 1")
    a.scale_list = [x.strip() for x in a.scales.split(",") if x.strip()]
    if not a.scale_list:
        ap.error("--scales 不能是空的")
    valid_scales = set(getattr(config, "SCALE_PROFILES", {}))
    bad = [s for s in a.scale_list if valid_scales and s not in valid_scales]
    if bad:
        ap.error(f"未知的規模 {bad}；可用：{sorted(valid_scales)}")
    a.config_list = [x.strip() for x in a.configs.split(",") if x.strip()]
    known = {name for name, _ in CONFIGS}
    bad = [c for c in a.config_list if c not in known]
    if bad:
        ap.error(f"未知的 config {bad}；可用：{sorted(known)}")
    if a.time_limit <= 0:
        ap.error("--time-limit 必須 > 0")
    if not (0.0 <= a.mip_gap < 1.0):
        ap.error("--mip-gap 必須在 [0, 1)")
    return a


def load_resume(path_str: str | None) -> dict[tuple, dict]:
    if not path_str:
        return {}
    p = Path(path_str)
    if not p.is_absolute():
        for cand in (Path.cwd() / p, RESULT_DIR / p.name):
            if cand.exists():
                p = cand
                break
    if not p.exists():
        print(f"[!] --resume-from 找不到 {p}，當成全新開始", flush=True)
        return {}
    done = {}
    with p.open(newline="", encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            if r.get("status") in SUCCESS_STATUS:
                for k in ("UB", "LB", "Gap_pct", "Time", "Nodes",
                          "cvar_tail_scenarios", "dro_relative_strength",
                          "root_seed_time", "root_seed_pct_of_limit",
                          "callback_pct_of_wall"):
                    try:
                        r[k] = float(r[k]) if r.get(k) not in (None, "") else None
                    except (TypeError, ValueError):
                        pass
                try:
                    r["n_scen"] = int(float(r["n_scen"]))
                except (TypeError, ValueError, KeyError):
                    continue
                done[(r.get("scale"), r["n_scen"], r.get("config"))] = r
    print(f"[續跑] 從 {p.name} 讀到 {len(done)} 格已完成的結果", flush=True)
    return done


def main(argv=None) -> int:
    a = parse_args(argv)
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    csv_path = RESULT_DIR / f"{RESULT_PREFIX}_raw_{stamp}.csv"
    xlsx_path = RESULT_DIR / f"{RESULT_PREFIX}_{stamp}.xlsx"

    cfg_by_name = dict(CONFIGS)
    cells = [(sc, n, cn) for sc in a.scale_list
             for n in a.scenario_list for cn in a.config_list]
    done = load_resume(a.resume_from)
    todo = [c for c in cells if c not in done]

    seed_budget = lshaped_core._resolve_root_seed_budget(a.time_limit)
    est = timedelta(seconds=len(todo) * a.time_limit)
    print(f"""
{'=' * 72}
  情境數掃描 pilot
{'=' * 72}
  規模          : {', '.join(a.scale_list)}
  情境數        : {', '.join(str(s) for s in a.scenario_list)}
  config        : {', '.join(a.config_list)}
  模型          : {a.model}{'（eps_mode=' + a.eps_mode + '）' if a.model == 'dro_box' else ''}
  每格時限      : {a.time_limit:,.0f} 秒 = {a.time_limit / 3600:.2f} 小時
  MIP gap 門檻  : {a.mip_gap}
  root seed 預算: {f'{seed_budget:,.0f} 秒（佔時限 {seed_budget / a.time_limit * 100:.0f}%）'
                   if seed_budget else '未設上限'}
  判定標準      : Extensive gap > {a.ext_min_gap:g}%　且　FullStack gap <= {a.fs_max_gap:g}%
  格數          : 共 {len(cells)} 格，待跑 {len(todo)} 格
  最壞總時長    : {est}（全部跑滿時限）
  輸出          : {xlsx_path.name}
{'=' * 72}
""", flush=True)
    if a.dry_run:
        print("  [DRY RUN] 不會呼叫 Gurobi，數字是假的，只驗流程。\n", flush=True)

    settings = {
        "產生時間": datetime.now().isoformat(timespec="seconds"),
        "用途": "掃描情境數 |S|，決定正式 ablation 用哪個 S",
        "規模": ", ".join(a.scale_list),
        "情境數": ", ".join(str(s) for s in a.scenario_list),
        "config": ", ".join(a.config_list),
        "模型": a.model,
        "eps_mode": a.eps_mode,
        "eps_bar×S 不變量": (EPS_BOX_PRODUCT if a.eps_mode == "scaled" else "（fixed 模式不適用）"),
        "每格時限(秒)": a.time_limit,
        "MIPGap": a.mip_gap,
        "判定_Extensive_gap下限(%)": a.ext_min_gap,
        "判定_FullStack_gap上限(%)": a.fs_max_gap,
        "BENDERS_ROOT_SEED_TIME_LIMIT": getattr(config, "BENDERS_ROOT_SEED_TIME_LIMIT", None),
        "root_seed 預算(秒)": seed_budget,
        "BENDERS_ROOT_SEED_ITERS": getattr(config, "BENDERS_ROOT_SEED_ITERS", None),
        "BENDERS_ROOT_CUT_ROUNDS": getattr(config, "BENDERS_ROOT_CUT_ROUNDS", None),
        "BENDERS_PARALLEL_ORACLES": getattr(config, "BENDERS_PARALLEL_ORACLES", None),
        "RISK_ALPHA": getattr(config, "RISK_ALPHA", None),
        "RISK_LAMBDA": getattr(config, "RISK_LAMBDA", None),
        "MASTER_SEED": getattr(config, "MASTER_SEED", None),
        "TIME_PERIODS": getattr(config, "TIME_PERIODS", None),
        "DISASTER_CSV": getattr(config, "DISASTER_CSV", None),
        "CCP_CSV": getattr(config, "CCP_CSV", None),
        "HOSPITAL_CSV": getattr(config, "HOSPITAL_CSV", None),
        "dry_run": a.dry_run,
        "續跑來源": a.resume_from or "（無）",
    }

    rows = [done[c] for c in cells if c in done]
    t_start = time.time()
    try:
        for i, (scale, n_scen, cname) in enumerate(todo, 1):
            print(f"\n[{i}/{len(todo)}] 已用 {timedelta(seconds=int(time.time() - t_start))}",
                  flush=True)
            rows.append(run_cell(scale, n_scen, cname, cfg_by_name[cname],
                                 a.model, a.eps_mode, a.time_limit, a.mip_gap,
                                 a.dry_run))
            rows.sort(key=lambda r: (r["scale"], int(r["n_scen"]), r["config"]))
            write_csv(csv_path, rows)                          # 每格寫一次
            try:
                write_xlsx(xlsx_path, rows,
                           build_verdict(rows, a.ext_min_gap, a.fs_max_gap), settings)
            except PermissionError:
                print(f"   [!] {xlsx_path.name} 正被開著，這次跳過 xlsx（CSV 已存）",
                      flush=True)
    except KeyboardInterrupt:
        print("\n\n[中斷] 已完成的結果都保住了。續跑指令：", flush=True)
        print(f'  python "{Path(__file__).name}" --resume-from "{csv_path.name}"',
              flush=True)

    verdict_rows = build_verdict(rows, a.ext_min_gap, a.fs_max_gap)
    write_csv(csv_path, rows)
    try:
        write_xlsx(xlsx_path, rows, verdict_rows, settings)
    except PermissionError:
        print(f"[!] {xlsx_path.name} 無法寫入（檔案開著？），CSV 仍在 {csv_path.name}",
              flush=True)

    print(f"\n{'=' * 72}\n  結果\n{'=' * 72}", flush=True)
    hdr = f"  {'scale':7} {'S':>4} {'Ext gap%':>10} {'FS gap%':>10} {'尾端情境':>9} {'判定':>6}"
    print(hdr)
    for v in verdict_rows:
        eg, fg = v["Extensive_gap_pct"], v["FullStack_gap_pct"]
        print(f"  {v['scale']:7} {v['n_scen']:>4} "
              f"{(f'{eg:.2f}' if isinstance(eg, (int, float)) else 'NA'):>10} "
              f"{(f'{fg:.2f}' if isinstance(fg, (int, float)) else 'NA'):>10} "
              f"{v['cvar_tail_scenarios']!s:>9} {v['verdict']:>6}"
              + (f"   {v['note']}" if v["note"] else ""))
    print(f"\n  {recommend(verdict_rows)}")
    print(f"\n  xlsx : {xlsx_path}")
    print(f"  csv  : {csv_path}")
    print(f"  總時長: {timedelta(seconds=int(time.time() - t_start))}\n", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
